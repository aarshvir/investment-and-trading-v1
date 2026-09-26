import json, math, hashlib
from pathlib import Path
from datetime import datetime,timezone
from openpyxl import load_workbook
P=Path(__file__).parent
j=json.loads((P/'decision_model.json').read_text())
wf=load_workbook(P/'Investment_Decision_Model.xlsx',data_only=False)
wv=load_workbook(P/'Investment_Decision_Model.xlsx',data_only=True)
checks=[];issues=[]
def check(label,a,b,tol=1e-8):
 ok=abs(a-b)<=tol
 checks.append({'check':label,'actual':a,'expected':b,'pass':ok})
 if not ok:issues.append(label)
def pv(ds,t,r):return sum(v/(1+r)**(i+1) for i,v in enumerate(ds))+t/(1+r)**3
def irr(p,ds,t):
 lo,hi=-.999,10
 for _ in range(160):
  m=(lo+hi)/2
  if pv(ds,t,m)>p:lo=m
  else:hi=m
 return (lo+hi)/2
for s in j['stocks']:
 for c in s['scenarios']:
  tag=s['ticker']+'/'+c['name'];ds=c['dividends'];t=c['terminal_eps']*c['terminal_pe'];r=s['hurdle']
  check(tag+'/terminal',c['terminal_price'],t)
  for suffix,p in [('reference',s['reference_price']),('limit',s['limit_price'])]:
   check(tag+'/irr_'+suffix,c['irr_'+suffix],irr(p,ds,t))
   check(tag+'/irr_'+suffix+'_withholding30',c['irr_'+suffix+'_withholding30'],irr(p,[d*.7 for d in ds],t))
   check(tag+'/return_'+suffix,c['total_return_'+suffix],(sum(ds)+t)/p-1)
  check(tag+'/pv',c['pv_hurdle'],pv(ds,t,r))
  check(tag+'/pv_withholding',c['pv_hurdle_withholding30'],pv([d*.7 for d in ds],t,r))
 check(s['ticker']+'/limit_within_tax_ceiling',float(s['limit_price']<=s['base_ceiling_withholding30']),1)
check('direct_initial',sum(s['initial_weight'] for s in j['stocks']),j['direct_initial'])
check('direct_full',sum(s['slot_weight'] for s in j['stocks']),j['direct_cap'])
check('initial_total',j['core_weight']+j['direct_initial']+j['reserve_initial'],1)
check('full_total',j['core_weight']+j['direct_cap']+j['reserve_min_full'],1)
check('direct_max_cap',float(max(s['slot_weight'] for s in j['stocks'])<=.02),1)
financial_weight=sum(s['slot_weight'] for s in j['stocks'] if s['sector']=='Financials')
check('direct_financial_cap',float(financial_weight<=.05),1)
for t in ('MSFT','NVDA'):
 wt=next(s['slot_weight'] for s in j['stocks'] if s['ticker']==t)+j['core_weight']*j['overlap'][t+'_index_weight']
 check(t+'/lookthrough',j['overlap'][t+'_total_full'],wt)
 check(t+'/lookthrough_cap',float(wt<=.035),1)
for p in j['portfolios']:
 cap=p['capital'];rows=p['rows'];check(str(cap)+'/total',sum(r['cost'] for r in rows),cap)
 for row in rows:
  check(str(cap)+'/'+row['ticker']+'/weight',row['weight'],row['cost']/cap)
  if row['whole_shares'] is not None:
   check(str(cap)+'/'+row['ticker']+'/cost',row['cost'],row['whole_shares']*row['illustrative_price'])
   check(str(cap)+'/'+row['ticker']+'/floor',row['whole_shares'],math.floor(row['budget']/row['illustrative_price']))
for s in j['stresses']:
 for prefix,dw,rw in [('initial',j['direct_initial'],j['reserve_initial']),('full',j['direct_cap'],j['reserve_min_full'])]:
  z=j['core_weight']*s['core_shock']+dw*s['direct_shock']+rw*s['reserve_shock'];check(s['name']+'/'+prefix,s[prefix+'_loss'],z)
 check(s['name']+'/50k',s['dollars_50k'],s['full_loss']*50000)
 check(s['name']+'/100k',s['dollars_100k'],s['full_loss']*100000)
for p in j['portfolio_scenarios']:
 for prefix,wn,rn in [('initial','initial_weight','reserve_initial'),('full','slot_weight','reserve_min_full')]:
  z=j['core_weight']*(1+p['index_annual_assumption'])**3+j[rn]*(1+p['reserve_annual_assumption'])**3
  for s in j['stocks']:
   c=next(c for c in s['scenarios'] if c['name']==p['name'])
   z+=s[wn]*(c['terminal_price']+.7*sum(c['dividends']))/s['limit_price']
  check(p['name']+'/'+prefix+'/wealth',p[prefix+'_terminal_wealth_multiple'],z)
  check(p['name']+'/'+prefix+'/cagr',p[prefix+'_wealth_cagr'],z**(1/3)-1)
formula_count=0;missing=0;excel_errors=[]
for ws in wf:
 for row in ws:
  for c in row:
   v=wv[ws.title][c.coordinate]
   if c.data_type=='f':
    formula_count+=1
    if v.value is None:missing+=1
   if v.data_type=='e':excel_errors.append(ws.title+'!'+c.coordinate+':'+str(v.value))
for row in range(2,44):
 ws=wv['Scenarios'];ticker=ws[f'A{row}'].value;case=ws[f'B{row}'].value
 st=next(s for s in j['stocks'] if s['ticker']==ticker);sc=next(c for c in st['scenarios'] if c['name']==case)
 vals={col:ws[f'{col}{row}'].value for col in ['C','D','E','F','G','H','I','J','K']}
 c,d,e,f,g,h,i,k,l=[vals[x] for x in ['C','D','E','F','G','H','I','J','K']]
 terminal=g*h;ds=[i,k,l]
 expected={'L':terminal,'M':irr(c,ds,terminal),'N':irr(d,ds,terminal),'O':irr(d,[v*(1-f) for v in ds],terminal),'P':pv(ds,terminal,e),'Q':pv([v*(1-f) for v in ds],terminal,e),'S':-c,'T':i,'U':k,'V':l+terminal,'W':-d,'X':i,'Y':k,'Z':l+terminal,'AA':-d,'AB':i*(1-f),'AC':k*(1-f),'AD':l*(1-f)+terminal}
 for col,v in expected.items():check('workbook/'+col+str(row),ws[f'{col}{row}'].value,v)
 for col,key in [('C','reference_price'),('D','limit_price'),('E','hurdle')]:check('workbook/inputs/'+ticker+'/'+case+'/'+col,vals[col],st[key])
 check('workbook/'+ticker+'/'+case+'/eps',g,sc['terminal_eps']);check('workbook/'+ticker+'/'+case+'/pe',h,sc['terminal_pe'])
for cap in (50000,100000):check('workbook/total/'+str(cap),wv[str(cap)+' initial']['F9'].value,cap)
cost_exceptions=[]
for s in j['stocks']:
 b=next(c for c in s['scenarios'] if c['name']=='base');ceiling=pv([v*.7 for v in b['dividends']],b['terminal_price']*.9975,s['hurdle'])/1.0025
 if s['limit_price']>ceiling:cost_exceptions.append({'ticker':s['ticker'],'limit':s['limit_price'],'ceiling_30pct_tax_25bp_each_side':ceiling})
out={'timestamp_utc':datetime.now(timezone.utc).isoformat(),'files':{n:hashlib.sha256((P/n).read_bytes()).hexdigest() for n in ['decision_model.json','Investment_Decision_Model.xlsx']},'numerical_checks':len(checks),'failed_checks':issues,'workbook_formula_count':formula_count,'missing_formula_cache':missing,'excel_errors':excel_errors,'cost_exceptions':cost_exceptions,'financial_direct_weight':financial_weight,'checks':checks}
(P/'consolidated_financials_audit.json').write_text(json.dumps(out,indent=2))
print({k:v for k,v in out.items() if k!='checks'})
