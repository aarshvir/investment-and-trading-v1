"""Reproducible Oct 4 stock-only weekly update from v014 and raw two-provider prices."""
import csv,hashlib,html,json,zipfile
from datetime import datetime,timezone
from pathlib import Path

H=Path(__file__).resolve().parent
R=H.parents[2]
P=["v014_2026-10-01_codex","v013_2026-10-01_codex","v012_2026-09-27_codex","v011_2026-09-27_claude"]
def read(p):return json.loads(p.read_text(encoding="utf-8"))
def write(name,obj):(H/name).write_text(json.dumps(obj,indent=2,ensure_ascii=False),encoding="utf-8")
def sha(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()

old=read(R/"GPT"/"2026-10-01"/"stocks_only_correction"/"STOCKS_ONLY_MODEL.json")
price=read(H/"agents"/"prices"/"PRICE_MANIFEST.json")
prior={x["ticker"]:x for x in old["records"]}
assert len(prior)==len(price["records"])==20
assert old["individual_stocks_pct"]==100 and old["treasury_bonds_pct"]==old["index_etf_pct"]==old["strategic_cash_pct"]==0
special={
"MCK":("WAIT_NEW_ADDITION_KEEP_TARGET","CVS agreement in principle; FY27 EPS reaffirmed, no definitive economics","2026-10-01 issuer release","2026-11-04 earnings and definitive terms","higher ~20.23x adjusted FY27 EPS entry"),
"PEP":("WAIT_NEW_ADDITION_KEEP_TARGET","H1 cash after capex below dividends","24-week FY2026 10-Q, inherited","2026-10-08 Q3 cash/dividend bridge","partial-year seasonality"),
"PTC":("WAIT_NEW_ADDITION_KEEP_TARGET","Oct1 price rise lacks verified earnings event","2026-10-02 SEC and issuer event sweep","next results date unverified","cause and valuation unverified"),
"AXP":("NO_CHANGE_KEEP_TARGET","GBTG sale; ~$975m projected pretax gain is nonrecurring","2026-10-01 13D/A; 2026-05-04 8-K","Q3 gain/tax/capital-use confirmation","actual gain/tax pending"),
"RSG":("NO_CHANGE_KEEP_TARGET","Oct2 $0.67 ex-dividend; do not count as received cash","2026-10-02 issuer and raw/adjusted quotes","next operating filing date unverified","holdings and dividend entitlement unknown")}
actions=[]; mark=0
for q in price["records"]:
    t=q["ticker"];v=prior[t]
    assert q["weight_pct"]==v["weight_pct"] and q["v014_sep30_usd"]==v["reference_close_usd"]
    assert q["identity_pass"] and q["price_match_penny"] and q["currency"]=="USD" and q["cboe_security_type"]=="stock"
    assert q["yahoo_rows"][-1]["date_et"]=="2026-10-02" and q["cboe_last_trade_time"].startswith("2026-10-02T")
    for provider in ("yahoo","cboe"):
        raw=H/"agents"/"prices"/"raw"/(provider+"_"+t+".json")
        assert raw.is_file() and sha(raw)==q[provider+"_raw_sha256"]
    move=q["oct2_close_usd"]/q["v014_sep30_usd"]-1
    mark+=move*q["weight_pct"]/100
    action,fact,source,next_event,uncertainty=special.get(t,("NO_CHANGE_KEEP_TARGET","No material operating filing identified in bounded sweep","2026-10-02 close and SEC submission index","next issuer result; date unverified","other issuer news may exist"))
    actions.append({"ticker":t,"prior_target_pct":q["weight_pct"],"revised_target_pct":q["weight_pct"],"action":action,
    "prior_standing_limit":"none","revised_standing_limit":"none; live/account gate","sep30_close_usd":q["v014_sep30_usd"],
    "oct2_close_usd":q["oct2_close_usd"],"move_pct":round(move*100,4),"changed_fact":fact,"source_period":source,
    "thesis_impact":"target unchanged; "+("new-addition gate" if t in ("MCK","PEP","PTC") else "monitor"),
    "next_catalyst":next_event,"allocation_effect_pct_points":0,"uncertainty":uncertainty})
assert sum(x["revised_target_pct"] for x in actions)==100
with (H/"ACTION_TABLE.csv").open("w",newline="",encoding="utf-8-sig") as f:
    w=csv.DictWriter(f,fieldnames=list(actions[0]));w.writeheader();w.writerows(actions)
scenarios=[
("Demand shock / sticky inflation",[-40,-25],[-8,1]),
("Continuing growth / valuation normalization",[-5,15],[4,8]),
("Strong productivity / AI cash payback",[15,30],[10,16])]
rows=[]
for name,one,annual in scenarios:
    rows.append({"state":name,"one_year_total_return_pct":one,"three_year_cagr_pct":annual,
    "three_year_total_return_pct":[round(((1+x/100)**3-1)*100,4) for x in annual],
    "usd_50000_after_three_years":[round(50000*(1+x/100)**3,2) for x in annual],
    "usd_100000_after_three_years":[round(100000*(1+x/100)**3,2) for x in annual]})
write("SCENARIO_MODEL.json",{"prepared_utc":datetime.now(timezone.utc).isoformat(),"parents":P,"price_cutoff_et":"2026-10-02",
"portfolio":"v014 100% individual stocks, weights unchanged","status":"judgmental conditional illustrations; no probabilities",
"basis":"nominal USD total return before investor-specific tax, FX and dealing costs",
"formula":"terminal wealth = initial capital*(1+annualized rate)^3","scenarios":rows,
"illustrative_tax_cost_only":{"assumed_gross_dividend_yield_pct":1.5,"assumption_is_audited_portfolio_yield":False,
"dividend_withholding_pct":30,"dividend_drag_pct_points_per_year":0.45,"purchase_and_sale_bps_each":25,
"rough_annual_three_year_dealing_drag_pct_points":0.17,"middle_band_after_these_hypothetical_drags_pct":[3.4,7.4]},
"v014_historical_gfc_path_drawdown_pct":-44.08,"price_only_mark_sep30_to_oct2_pct":round(mark*100,4)})
write("SOURCE_LEDGER.json",{"prepared_utc":datetime.now(timezone.utc).isoformat(),
"price_manifest":"agents/prices/PRICE_MANIFEST.json","raw_quote_count":40,
"sec_sweep":{"ticker_count":20,"through":"2026-10-02","new_operating_8k_10q_10k":0,"limit":"issuer releases outside SEC can be material","MCK_submission_sha256":"d9a035c9346942d3dee57e82b1903131799f0d45ef16bc179bd40d50a82073fa","PTC_submission_sha256":"77680f19c4f505ce908009f732c4a9c9296252cc0117cb75555c04bf4fede6ac"},
"sources":[
{"fact":"September jobs +29k; unemployment 4.2%; revisions -60k","period":"2026-09; released 2026-10-02","url":"https://www.bls.gov/news.release/archives/empsit_10022026.htm"},
{"fact":"10-year par yield 5.28%","period":"2026-10-02","url":"https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?field_tdr_date_value=2026&type=daily_treasury_yield_curve"},
{"fact":"August PCE headline/core 3.4%/3.0%","period":"2026-08","url":"https://www.bea.gov/news/2026/personal-income-and-outlays-august-2026"},
{"fact":"MCK CVS agreement in principle; FY27 adjusted EPS $44.20-$45.00 reaffirmed","period":"issuer release 2026-10-01","url":"https://www.businesswire.com/news/home/20261001569834/en/McKesson-Corporation-Extends-Pharmaceutical-Distribution-Agreement-with-CVS-Health","syndication_sha256":"03a1d2ceee53ad5a1e0ac330702e3227af0fdd50fe9ceb9fe7f9ec10e416929e"},
{"fact":"CVS ~24% revenue/~21% receivables","period":"MCK FY2026 10-K","url":"https://www.sec.gov/Archives/edgar/data/927653/000092765326000069/mck-20260331.htm","raw_sha256":"69994a40880089f06b9da5fa18e4cc9966cec1163876d12492a63fcce8740b42"},
{"fact":"AXP GBTG sale close","period":"2026-10-01 13D/A; economic close 2026-09-29","url":"https://www.sec.gov/Archives/edgar/data/4962/000000496226000356/primary_doc.xml","raw_sha256":"342317448c313d9819574bb93e7f6eb0467867175aba9637831a733d6d919307"},
{"fact":"AXP projected ~$975m one-time pretax gain, excluded from guidance","period":"2026-05-04 8-K","url":"https://www.sec.gov/Archives/edgar/data/4962/000000496226000203/axp-20260504.htm"},
{"fact":"PEP Q3 results scheduled Oct8","period":"issuer release","url":"https://www.pepsico.com/en/newsroom/press-releases/2026/pepsico-announces-timing-and-availability-of-third-quarter-2026-financial-results"}]})
write("VALIDATION.json",{"status":"pass_with_limits","checked_utc":datetime.now(timezone.utc).isoformat(),"ticker_count":20,
"target_weight_pct":100,"nonstock_weight_pct":0,"dual_source_identity_count":20,"raw_quote_hashes_rechecked":40,
"price_only_mark_pct":round(mark*100,4),"allocation_change_pct_points":0,"broader_universe_refreshed":False,
"actual_holdings_known":False,"full_comparable_company_forecasts":False})
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
wb=Workbook();ws=wb.active;ws.title="Weekly actions"
ws.append(["Ticker","Target %","Sep30","Oct2","Move %","Action","Fact","Next catalyst"])
for a in actions:ws.append([a["ticker"],a["revised_target_pct"],a["sep30_close_usd"],a["oct2_close_usd"],a["move_pct"],a["action"],a["changed_fact"],a["next_catalyst"]])
ss=wb.create_sheet("Forward states");ss.append(["State","1y low %","1y high %","3y CAGR low %","3y CAGR high %","$50k low","$50k high","$100k low","$100k high"])
for x in rows:ss.append([x["state"],*x["one_year_total_return_pct"],*x["three_year_cagr_pct"],*x["usd_50000_after_three_years"],*x["usd_100000_after_three_years"]])
for sheet in wb:
    sheet.freeze_panes="A2"
    for cell in sheet[1]:cell.fill=PatternFill("solid",fgColor="15304A");cell.font=Font(color="FFFFFF",bold=True)
    for col in sheet.columns:sheet.column_dimensions[get_column_letter(col[0].column)].width=min(65,max(12,max(len(str(c.value or "")) for c in col)+2))
wb.save(H/"WEEKLY_STOCKS_WORKBOOK.xlsx")
p_rows="".join("<tr><td>{}</td><td>{}%</td><td>USD {:,.2f}</td><td>{:+.2f}%</td><td>{}</td></tr>".format(html.escape(a["ticker"]),a["revised_target_pct"],a["oct2_close_usd"],a["move_pct"],html.escape(a["action"])) for a in actions)
s_rows="".join("<tr><td>{}</td><td>{:+}% to {:+}%</td><td>{:+}% to {:+}%</td><td>USD {:,.0f} to {:,.0f}</td></tr>".format(html.escape(x["state"]),*x["one_year_total_return_pct"],*x["three_year_cagr_pct"],*x["usd_50000_after_three_years"]) for x in rows)
dashboard='<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Stocks-only weekly review</title><style>body{background:#0c1626;color:#e8edf4;font:16px/1.6 system-ui;margin:0}main{max-width:1150px;margin:auto;padding:35px 23px}a{color:#bddcff}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:12px}.card{background:#15283e;padding:18px;border:1px solid #35506b;border-radius:8px}.card strong{display:block;font-size:27px}.warn{background:#33291e;border-left:4px solid #e9b76c;padding:15px}table{border-collapse:collapse;width:100%;background:#15283e}td,th{padding:9px;border-bottom:1px solid #344a62;text-align:left}th{background:#1f3853}.overflow{overflow-x:auto}</style><main><p>CODEX · OCT 2 COMPLETED PRICES</p><h1>Stocks-only: no weight change</h1><div class="cards"><div class="card">Target<strong>100% stocks</strong></div><div class="card">Price-only mark<strong>+{:.2f}%</strong></div><div class="card">3y middle scenario<strong>4–8%/yr</strong></div><div class="card">Historical GFC path<strong>−44.08%</strong></div></div><p class="warn">Illustrative ranges, no assigned odds. A 100% stock portfolio can exceed the previously stated 15–20% tolerable decline.</p><h2>Future states</h2><div class="overflow"><table><tr><th>State</th><th>1-year</th><th>3-year CAGR</th><th>$50k after 3 years</th></tr>{}</table></div><h2>Weekly target actions</h2><div class="overflow"><table><tr><th>Stock</th><th>Target</th><th>Oct2 close</th><th>Move</th><th>New-money gate</th></tr>{}</table></div><p>See the decision report for source periods, tax/cost assumptions, and MCK/PEP/PTC evidence gates. No actual holdings or trades are asserted.</p></main></html>'.replace('{:.2f}',f'{mark*100:.2f}').replace('{}',s_rows,1).replace('{}',p_rows,1)
(H/"WEEKLY_DASHBOARD.html").write_text(dashboard,encoding="utf-8")
package=H/"GPT_Weekly_Stocks_2026-10-04.zip"
files=[p for p in H.rglob("*") if p.is_file() and p!=package and p.name!="PACKAGE_MANIFEST.json"]
write("PACKAGE_MANIFEST.json",{"prepared_utc":datetime.now(timezone.utc).isoformat(),"files":[{"path":p.relative_to(H).as_posix(),"bytes":p.stat().st_size,"sha256":sha(p)} for p in sorted(files)]})
files.append(H/"PACKAGE_MANIFEST.json")
with zipfile.ZipFile(package,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in sorted(files):z.write(p,"GPT/2026-10-04/weekly/"+p.relative_to(H).as_posix())
    z.write(R/"versions"/P[0]/"package.zip","parent_completed/v014_package.zip")
    for name in ("WEEKLY_DECISION_AND_SCENARIOS.md","ACTION_TABLE.csv","SCENARIO_MODEL.json","WEEKLY_STOCKS_WORKBOOK.xlsx"):
        z.write(H/name,"review/weekly/2026-10-04_"+name)
with zipfile.ZipFile(package) as z:assert z.testzip() is None
print(json.dumps({"package":str(package),"bytes":package.stat().st_size,"sha256":sha(package),"price_only_mark_pct":round(mark*100,4),"files":len(files)}))

